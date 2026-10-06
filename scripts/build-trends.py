#!/usr/bin/env python3
"""Build the combined enrolment, pupil-characteristic, and attendance data."""

from __future__ import annotations

import json
import importlib.util
import re
from datetime import date
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote, urljoin
from urllib.request import Request, urlopen

from openpyxl import load_workbook

SOURCE_ROOT = "https://www.education-ni.gov.uk"
ATTENDANCE_INDEX = f"{SOURCE_ROOT}/articles/pupil-attendance"
CENSUS_INFOGRAPHICS = f"{SOURCE_ROOT}/articles/education-data-infographics"
DATA_DIR = Path(__file__).resolve().parents[1] / "data"
OUTPUT_PATH = DATA_DIR / "school-trends.json"
ENROLMENT_MODULE_PATH = Path(__file__).with_name("build-enrolments.py")
YEAR_PATTERN = re.compile(r"(?<!\d)(20\d{2})[/_-](\d{2})(?!\d)")

_enrolment_spec = importlib.util.spec_from_file_location("build_enrolments", ENROLMENT_MODULE_PATH)
if _enrolment_spec is None or _enrolment_spec.loader is None:
    raise RuntimeError(f"Could not load enrolment importer at {ENROLMENT_MODULE_PATH}.")
_enrolment_module = importlib.util.module_from_spec(_enrolment_spec)
_enrolment_spec.loader.exec_module(_enrolment_module)

DEMOGRAPHIC_FILES = {
    "irish-medium": {
        "label": "Irish-medium education",
        "title": "Pupils in Irish-medium education",
        "unit": "pupils",
        "format": "count",
        "filename": "irish-medium.xlsx",
        "match": "children in irish medium education",
        "sheet": "pupils",
        "notes": ["Includes pupils in funded Irish-medium schools and units in English-medium schools."],
    },
    "free-meals": {
        "label": "Free school meal entitlement",
        "title": "Percentage entitled to free school meals",
        "unit": "%",
        "format": "percent",
        "filename": "free-meals.xlsx",
        "match": "percentage pupils entitled to free school meals",
        "sheet": "percent fsme",
        "notes": ["Values are the published percentage of pupils entitled to free school meals; they are not counts."],
    },
    "newcomers": {
        "label": "Newcomer pupils",
        "title": "Newcomer pupil counts",
        "unit": "pupils",
        "format": "count",
        "filename": "newcomers.xlsx",
        "match": "newcomer pupils",
        "sheet": "newcomers",
        "notes": [
            "A newcomer pupil is defined by the Department as a pupil who lacks the language skills to participate fully in the school curriculum and wider environment.",
            "Values suppressed for disclosure or marked as fewer than five are shown as unavailable.",
        ],
    },
}


class AnchorParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[tuple[str, str]] = []
        self._href: str | None = None
        self._text: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "a":
            self._href = dict(attrs).get("href")
            self._text = []

    def handle_data(self, data: str) -> None:
        if self._href is not None:
            self._text.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "a" and self._href:
            text = " ".join(" ".join(self._text).split())
            self.links.append((urljoin(SOURCE_ROOT, self._href), text))
            self._href = None
            self._text = []


def fetch(url: str) -> bytes:
    request = Request(url, headers={"User-Agent": "NorthernIrelandSchools data importer"})
    with urlopen(request, timeout=60) as response:
        return response.read()


def anchors_at(url: str) -> list[tuple[str, str]]:
    parser = AnchorParser()
    parser.feed(fetch(url).decode("utf-8", errors="replace"))
    return parser.links


def academic_year(value: object) -> tuple[int, str] | None:
    match = YEAR_PATTERN.search(str(value or ""))
    if not match:
        return None
    start_year = int(match.group(1))
    return start_year, f"{start_year}/{match.group(2)}"


def number(value: object) -> float | int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else float(value)
    text = str(value).strip().replace(",", "")
    try:
        result = float(text)
    except ValueError:
        return None
    return int(result) if result.is_integer() else result


def readable_label(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\ufffd", "—").strip())


def row_id(label: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", label.casefold()).strip("-")
    return slug or "series"


def source_links() -> tuple[dict[str, str], str]:
    links = anchors_at(_enrolment_module.SOURCE_PAGE)
    sources: dict[str, str] = {}
    for key, config in DEMOGRAPHIC_FILES.items():
        source = next(
            (
                href
                for href, text in links
                if config["match"] in unquote(href).casefold()
                or config["match"] in text.casefold()
            ),
            None,
        )
        if not source:
            raise RuntimeError(f"Could not find the Department's {config['label']} workbook.")
        sources[key] = source
    religion_source = next(
        (
            href
            for href, text in links
            if "pupil religion by school management type" in unquote(href).casefold()
            or "pupil religion by school management type" in text.casefold()
        ),
        None,
    )
    if not religion_source:
        raise RuntimeError("Could not find the Department's pupil religion workbook.")
    sources["religion"] = religion_source
    return sources, _enrolment_module.SOURCE_PAGE


def find_year_columns(rows: list[tuple[object, ...]]) -> tuple[int, list[tuple[int, str]]]:
    for row_index, row in enumerate(rows):
        year_columns = [
            (column_index, parsed[1])
            for column_index, value in enumerate(row)
            if (parsed := academic_year(value)) is not None
        ]
        if len(year_columns) >= 5:
            return row_index, year_columns
    raise RuntimeError("Could not find an academic-year header in a demographic workbook.")


def build_wide_demographic_dataset(
    key: str,
    config: dict[str, str | list[str]],
    source_url: str,
) -> dict[str, object]:
    workbook_path = DATA_DIR / config["filename"]
    workbook_bytes = fetch(source_url)
    workbook_path.write_bytes(workbook_bytes)
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    sheet_name = str(config["sheet"])
    if sheet_name not in workbook.sheetnames:
        raise RuntimeError(f"The {config['label']} workbook is missing its {sheet_name!r} sheet.")
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    header_index, year_columns = find_year_columns(rows)
    years = [
        {"startYear": academic_year(label)[0], "label": label}
        for _, label in year_columns
    ]
    series: list[dict[str, object]] = []
    active_group = ""
    used_ids: set[str] = set()

    for row in rows[header_index + 1 :]:
        label = readable_label(row[0] if row else None)
        if not label:
            continue
        raw_values = [row[column] if column < len(row) else None for column, _ in year_columns]
        parsed_values = [number(value) for value in raw_values]
        has_reported_value = any(
            value is not None and str(value).strip().casefold() not in {"", "-", "n.a.", "na"}
            for value in raw_values
        )
        if not has_reported_value:
            if not label.casefold().startswith(("note", "figure", "n.a", "source", "1.", "2.", "3.")):
                active_group = label
            continue
        if label.casefold().startswith(("note", "figure", "source", "n.a", "1.", "2.", "3.")):
            continue

        display_label = label
        if label.casefold() in {"total", "full-time", "part-time", "schools", "units"} and active_group:
            display_label = f"{active_group} — {label}"
        if key == "free-meals":
            parsed_values = [
                round(float(value) * 100, 4) if value is not None else None
                for value in parsed_values
            ]
        values = {
            year_label: parsed_values[index]
            for index, (_, year_label) in enumerate(year_columns)
        }
        if not any(value is not None for value in values.values()):
            continue
        base_id = row_id(display_label)
        identifier = base_id
        suffix = 2
        while identifier in used_ids:
            identifier = f"{base_id}-{suffix}"
            suffix += 1
        used_ids.add(identifier)
        series.append({
            "id": identifier,
            "label": display_label,
            "group": active_group or None,
            "values": values,
        })

    workbook.close()
    if not series:
        raise RuntimeError(f"The {config['label']} workbook contained no time-series rows.")
    return {
        "id": key,
        "label": config["label"],
        "title": config["title"],
        "unit": config["unit"],
        "format": config["format"],
        "years": years,
        "rows": series,
        "sourceUrl": source_url,
        "sourceFile": f"./data/{config['filename']}",
        "relatedLinks": [{"label": "Pupil data infographics", "url": CENSUS_INFOGRAPHICS}],
        "notes": config["notes"],
    }


def religion_categories(header: tuple[object, ...]) -> dict[str, list[int]]:
    columns: dict[str, list[int]] = {"Protestant": [], "Catholic": [], "Other / not recorded": []}
    for index, value in enumerate(header):
        label = readable_label(value).casefold()
        if "protestant" in label:
            columns["Protestant"].append(index)
        elif "catholic" in label:
            columns["Catholic"].append(index)
        elif any(term in label for term in ("other", "non christian", "no religion", "not recorded")):
            columns["Other / not recorded"].append(index)
    return {key: indices for key, indices in columns.items() if indices}


def sum_cells(row: tuple[object, ...], columns: list[int]) -> int | None:
    values = [number(row[index]) if index < len(row) else None for index in columns]
    if not values or any(value is None for value in values):
        return None
    return int(sum(values))


def build_religion_dataset(source_url: str) -> dict[str, object]:
    filename = "religion.xlsx"
    workbook_path = DATA_DIR / filename
    workbook_bytes = fetch(source_url)
    workbook_path.write_bytes(workbook_bytes)
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    values: dict[str, dict[str, int | None]] = {}
    groups: dict[str, str] = {}
    labels: dict[str, str] = {}
    years_by_start: dict[int, str] = {}

    for sheet in workbook.worksheets:
        sheet_year = academic_year(sheet.title.replace("_", "/"))
        if not sheet_year:
            continue
        start_year, year_label = sheet_year
        rows = list(sheet.iter_rows(values_only=True))
        header_index = next(
            (
                index
                for index, row in enumerate(rows)
                if sum(1 for value in row if "protestant" in readable_label(value).casefold()) > 0
            ),
            None,
        )
        if header_index is None:
            continue
        categories = religion_categories(rows[header_index])
        if not categories:
            continue
        years_by_start[start_year] = year_label
        active_group = ""

        for row in rows[header_index + 1 :]:
            label = readable_label(row[0] if row else None)
            if not label:
                continue
            relevant = [
                number(row[index]) if index < len(row) else None
                for indices in categories.values()
                for index in indices
            ]
            if all(value is None for value in relevant):
                if label.casefold().startswith("all funded pre-school"):
                    active_group = "All funded pupils"
                elif label.casefold().startswith(("note", "figure", "source", "*", "#", "1.", "2.", "3.")):
                    continue
                else:
                    active_group = label
                continue

            if label.casefold() == "total":
                group = active_group or "Unclassified"
            elif label.casefold().startswith("all funded pre-school"):
                group = "All funded pupils"
            else:
                continue
            for category, columns in categories.items():
                identifier_label = f"{group} — {category}"
                identifier = row_id(identifier_label)
                values.setdefault(identifier, {})[year_label] = sum_cells(row, columns)
                groups[identifier] = group
                labels[identifier] = identifier_label

    workbook.close()
    years = [
        {"startYear": start_year, "label": year_label}
        for start_year, year_label in sorted(years_by_start.items())
    ]
    all_year_labels = [year["label"] for year in years]
    religion_rows = [
        {
            "id": identifier,
            "label": labels[identifier],
            "group": groups[identifier],
            "values": {year: values[identifier].get(year) for year in all_year_labels},
        }
        for identifier in values
    ]
    if not religion_rows:
        raise RuntimeError("The pupil religion workbook contained no usable annual totals.")
    return {
        "id": "religion",
        "label": "Religion",
        "title": "Pupils by religion and school type",
        "unit": "pupils",
        "format": "count",
        "years": years,
        "rows": religion_rows,
        "sourceUrl": source_url,
        "sourceFile": f"./data/{filename}",
        "relatedLinks": [{"label": "Pupil data infographics", "url": CENSUS_INFOGRAPHICS}],
        "notes": [
            "The Department groups religion as Protestant, Catholic, and Other Christian / non-Christian / no religion / not recorded in recent years.",
            "Rows are school-type subtotals, not mutually exclusive demographic samples; use the school-type subtotals for comparisons.",
            "Small counts suppressed in the source are shown as unavailable. The Department did not publish a 2023/24 sheet in this workbook.",
        ],
    }


def attendance_pages() -> list[tuple[int, str, str]]:
    pages: dict[int, tuple[str, str]] = {}
    for href, text in anchors_at(ATTENDANCE_INDEX):
        if "pupil attendance" not in text.casefold():
            continue
        parsed = academic_year(text)
        if parsed:
            pages[parsed[0]] = (parsed[1], href)
    return [(start_year, label, href) for start_year, (label, href) in sorted(pages.items())]


def attendance_workbook_url(publication_url: str) -> str | None:
    for href, _ in anchors_at(publication_url):
        decoded = unquote(href).casefold()
        if ".xlsx" in decoded and "attendance" in decoded:
            return href
    return None


def attendance_measure(header: tuple[object, ...]) -> dict[str, int]:
    columns: dict[str, int] = {}
    for index, value in enumerate(header):
        label = readable_label(value).casefold().replace("authorised", "authorized")
        if "attendance" in label:
            columns["attendance"] = index
        elif "unauthorized absence" in label or "unauthorised absence" in label:
            columns["unauthorized-absence"] = index
        elif "authorized absence" in label or "authorised absence" in label:
            columns["authorized-absence"] = index
        elif label.strip() == "% absence" or label.strip() == "absence":
            columns["absence"] = index
    return columns


def attendance_percentage(value: object) -> float | None:
    parsed = number(value)
    if parsed is None or parsed < 0 or parsed > 100:
        return None
    return round(float(parsed), 2)


def build_attendance_dataset() -> dict[str, object]:
    workbooks: list[tuple[int, str, str, Path]] = []
    for start_year, year_label, page_url in attendance_pages():
        workbook_url = attendance_workbook_url(page_url)
        if not workbook_url:
            continue
        file_path = DATA_DIR / f"attendance-{year_label.replace('/', '-')}.xlsx"
        file_path.write_bytes(fetch(workbook_url))
        workbooks.append((start_year, year_label, workbook_url, file_path))

    if not workbooks:
        raise RuntimeError("Could not find downloadable annual attendance workbooks.")

    directory_path = DATA_DIR / "schools.json"
    directory_records = json.loads(directory_path.read_text(encoding="utf-8"))["schools"]
    directory = {
        re.sub(r"\D", "", str(school["id"])): school
        for school in directory_records
    }
    school_rows: dict[str, dict[str, object]] = {}
    for _, year_label, _, file_path in workbooks:
        workbook = load_workbook(file_path, data_only=True, read_only=True)
        for sheet in workbook.worksheets:
            if sheet.title.casefold() == "main menu":
                continue
            sheet_rows = sheet.iter_rows(values_only=True)
            header = None
            reference_column = None
            name_column = None
            metric_columns: dict[str, int] = {}
            for candidate in sheet_rows:
                if not candidate:
                    continue
                header_labels = {
                    readable_label(value).casefold(): index
                    for index, value in enumerate(candidate)
                    if value
                }
                candidate_reference_column = next(
                    (
                        index
                        for label, index in header_labels.items()
                        if "reference number" in label or "deni no" in label or "ref no" in label
                    ),
                    None,
                )
                candidate_metric_columns = attendance_measure(candidate)
                if candidate_reference_column is None or not candidate_metric_columns:
                    continue
                header = candidate
                reference_column = candidate_reference_column
                name_column = next(
                    (index for label, index in header_labels.items() if "school name" in label),
                    None,
                )
                metric_columns = candidate_metric_columns
                break
            if header is None or reference_column is None or not metric_columns:
                continue
            phase = sheet.title
            if phase.casefold().startswith("primary"):
                phase = "Primary"
            elif phase.casefold().startswith("post"):
                phase = "Post-primary"
            elif phase.casefold().startswith("special"):
                phase = "Special"
            for row in sheet_rows:
                raw_identifier = row[reference_column] if reference_column < len(row) else None
                identifier = re.sub(r"\D", "", str(raw_identifier or ""))
                if not identifier:
                    continue
                raw_name = row[name_column] if name_column is not None and name_column < len(row) else identifier
                school = directory.get(identifier)
                entry = school_rows.setdefault(identifier, {
                    "id": identifier,
                    "label": school["name"] if school else readable_label(raw_name),
                    "group": phase,
                    "values": {},
                })
                entry["values"][year_label] = {
                    metric: attendance_percentage(row[column] if column < len(row) else None)
                    for metric, column in metric_columns.items()
                }
        workbook.close()

    years = [
        {"startYear": start_year, "label": year_label}
        for start_year, year_label, _, _ in workbooks
    ]
    rows = sorted(
        school_rows.values(),
        key=lambda entry: (str(entry["group"]), str(entry["label"]).casefold()),
    )
    return {
        "id": "attendance",
        "label": "Attendance by school",
        "title": "School attendance and absence rates",
        "unit": "%",
        "format": "percent",
        "metricOptions": [
            {"id": "attendance", "label": "Attendance"},
            {"id": "absence", "label": "Absence"},
            {"id": "authorized-absence", "label": "Authorised absence"},
            {"id": "unauthorized-absence", "label": "Unauthorised absence"},
        ],
        "defaultMetric": "attendance",
        "years": years,
        "rows": rows,
        "sourceUrl": ATTENDANCE_INDEX,
        "sourceFile": f"./data/attendance-{years[-1]['label'].replace('/', '-')}.xlsx",
        "notes": [
            "Attendance data is published as a rate for each grant-aided school, not a pupil count. Rows are individual schools.",
            "Annual school-level spreadsheets are available from 2007/08; no downloadable school-level workbook is linked for 2022/23.",
            "School attendance rates should not be averaged to estimate a Northern Ireland-wide rate; the Department's pupil-weighted overall statistics are published separately.",
            "The annual attendance bulletins also discuss sector and pupil-characteristic breakdowns such as gender and ethnicity; the downloadable workbooks used here contain school-level rates.",
        ],
    }


def main() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    enrolment_url = _enrolment_module.find_workbook_url()
    enrolment_bytes = fetch(enrolment_url)
    _enrolment_module.WORKBOOK_PATH.write_bytes(enrolment_bytes)
    enrolment_data = _enrolment_module.build_data(enrolment_bytes, enrolment_url)
    sources, _ = source_links()
    datasets = [enrolment_dataset(enrolment_data)]
    datasets.append(build_religion_dataset(sources["religion"]))
    for key, config in DEMOGRAPHIC_FILES.items():
        datasets.append(build_wide_demographic_dataset(key, config, sources[key]))
    datasets.append(build_attendance_dataset())

    combined = {
        "snapshotDate": date.today().isoformat(),
        "source": "Department of Education Northern Ireland",
        "datasets": datasets,
    }
    OUTPUT_PATH.write_text(
        json.dumps(combined, ensure_ascii=False, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )
    print(
        f"Wrote {len(datasets)} datasets and "
        f"{sum(len(item['rows']) for item in datasets)} series to {OUTPUT_PATH}"
    )


def enrolment_dataset(data: dict[str, object]) -> dict[str, object]:
    rows = [
        {
            "id": category["id"],
            "label": category["label"],
            "group": None,
            "values": category["values"],
        }
        for category in data["categories"]
    ]
    return {
        "id": "school-category",
        "label": "School category",
        "title": "Pupils enrolled by school category",
        "unit": "pupils or funded places",
        "format": "count",
        "years": data["years"],
        "rows": rows,
        "sourceUrl": data["sourcePage"],
        "sourceFile": "./data/enrolment-by-school-management-type.xlsx",
        "relatedLinks": [{"label": "Pupil data infographics", "url": CENSUS_INFOGRAPHICS}],
        "notes": data["notes"] + [
            "The Department also publishes annual infographics on pupil characteristics such as sex, age, ethnicity and special educational needs; they are not supplied as a consistent time-series workbook here."
        ],
    }


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Download the Department's enrolment workbook and create the site data file."""

from __future__ import annotations

import json
import re
from datetime import date
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote, urljoin
from urllib.request import Request, urlopen

from openpyxl import load_workbook


SOURCE_PAGE = "https://www.education-ni.gov.uk/publications/school-enrolments-northern-ireland-summary-data"
SOURCE_ROOT = "https://www.education-ni.gov.uk"
DATA_DIR = Path(__file__).resolve().parents[1] / "data"
WORKBOOK_PATH = DATA_DIR / "enrolment-by-school-management-type.xlsx"
JSON_PATH = DATA_DIR / "enrolment-trends.json"

CATEGORIES = (
    ("funded-preschool", "Funded pre-school places"),
    ("nursery", "Nursery"),
    ("specialist-preschool", "Pre-school specialist provision"),
    ("primary", "Primary"),
    ("secondary", "Secondary"),
    ("special", "Special"),
    ("eotas", "EOTAS"),
    ("independent", "Independent"),
    ("hospital", "Hospital"),
    ("total", "Grand total incl. funded places"),
)


class LinkParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "a":
            href = dict(attrs).get("href")
            if href:
                self.links.append(urljoin(SOURCE_ROOT, href))


def fetch(url: str) -> bytes:
    request = Request(url, headers={"User-Agent": "NorthernIrelandSchools data importer"})
    with urlopen(request, timeout=60) as response:
        return response.read()


def find_workbook_url() -> str:
    parser = LinkParser()
    parser.feed(fetch(SOURCE_PAGE).decode("utf-8", errors="replace"))
    for url in parser.links:
        filename = unquote(url.rsplit("/", 1)[-1]).casefold()
        if filename.endswith(".xlsx") and "enrolment by school management type" in filename:
            return url
    raise RuntimeError("The Department's summary-data page did not link its enrolment workbook.")


def numeric_value(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    text = str(value or "").strip().replace(",", "")
    if text.isdigit():
        return int(text)
    return None


def year_info(sheet_name: str) -> tuple[int, str] | None:
    match = re.search(r"(\d{4})[_-](\d{2})$", sheet_name)
    if not match:
        return None
    start_year = int(match.group(1))
    return start_year, f"{start_year}/{str(start_year + 1)[-2:]}"


def normalized_label(value: object) -> str:
    label = re.sub(r"\s*\d+$", "", str(value or "").strip())
    return re.sub(r"\s+", " ", label).upper()


def extract_sheet(sheet: object, start_year: int) -> dict[str, int | None]:
    rows = list(sheet.iter_rows(values_only=True))
    header = next(
        (row for row in rows if any(str(value or "").strip().casefold() == "grand total" for value in row)),
        None,
    )
    if header is None:
        raise RuntimeError(f"Sheet {sheet.title!r} has no Grand Total column.")
    grand_total_column = next(
        index for index, value in enumerate(header)
        if str(value or "").strip().casefold() == "grand total"
    )
    preschool_columns = [
        index for index, value in enumerate(header)
        if "pre-school centre" in str(value or "").strip().casefold()
    ]

    values: dict[str, int | None] = {}
    active_section: str | None = None
    completed_sections: set[str] = set()

    def save(category_id: str, row: tuple[object, ...], exclude_funded_places: bool = False) -> None:
        value = row[grand_total_column] if grand_total_column < len(row) else None
        count = numeric_value(value)
        if exclude_funded_places and count is not None:
            overlap = sum(
                numeric_value(row[index]) or 0
                for index in preschool_columns
                if index < len(row)
            )
            count -= overlap
        values[category_id] = count

    for row in rows:
        label = normalized_label(row[0] if row else None)
        if not label:
            continue
        if label.startswith("VOLUNTARY AND PRIVATE PRESCHOOLS"):
            save("funded-preschool", row)
            completed_sections.add("funded-preschool")
            active_section = None
            continue
        if label.startswith("PRE-SCHOOL EDUCATION"):
            active_section = "funded-preschool"
            continue
        if label == "NURSERY":
            active_section = "nursery"
            continue
        if label == "PRIMARY":
            active_section = "primary"
            continue
        if label.startswith("SECONDARY"):
            active_section = "secondary"
            continue
        if label.startswith("PRE-SCHOOL SPECIALIST PROVISION"):
            active_section = "specialist-preschool"
            continue
        if label.startswith("GRAND TOTAL"):
            save("total", row)
            continue
        if label.startswith("TOTAL"):
            if active_section and active_section not in completed_sections:
                save(active_section, row, exclude_funded_places=active_section == "nursery")
                completed_sections.add(active_section)
            continue
        if label == "SPECIAL":
            save("special", row)
        elif label == "HOSPITAL":
            save("hospital", row)
        elif label == "INDEPENDENT":
            save("independent", row)
        elif label.startswith("EDUCATED OFF SITE BUT NOT AT SCHOOL"):
            save("eotas", row)

    if "total" not in values:
        raise RuntimeError(f"Sheet {sheet.title!r} has no Grand Total pupil count.")
    if "funded-preschool" not in values:
        values["funded-preschool"] = None
    return values


def build_data(workbook_bytes: bytes, workbook_url: str) -> dict[str, object]:
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    annual_data: list[tuple[int, str, dict[str, int | None]]] = []
    for sheet in workbook.worksheets:
        parsed_year = year_info(sheet.title)
        if parsed_year:
            start_year, year_label = parsed_year
            annual_data.append((start_year, year_label, extract_sheet(sheet, start_year)))
    workbook.close()
    annual_data.sort(key=lambda item: item[0])
    if not annual_data:
        raise RuntimeError("No annual enrolment data sheets were found in the workbook.")

    years = [{"startYear": start_year, "label": label} for start_year, label, _ in annual_data]
    categories = [
        {
            "id": category_id,
            "label": label,
            "values": {
                year_label: values.get(category_id)
                for _, year_label, values in annual_data
            },
        }
        for category_id, label in CATEGORIES
        if any(category_id in values for _, _, values in annual_data)
    ]
    return {
        "snapshotDate": date.today().isoformat(),
        "source": "Department of Education Northern Ireland",
        "sourcePage": SOURCE_PAGE,
        "sourceWorkbook": workbook_url,
        "measure": "School enrolments and funded pre-school places",
        "geography": "Northern Ireland",
        "years": years,
        "categories": categories,
        "notes": [
            "This is a Northern Ireland-level summary, not school-by-school data.",
            "The source Grand Total includes school enrolments and funded pre-school places. Funded places are shown separately; where the source groups them into the Nursery total, they are split out to avoid double-counting.",
            "A dash or blank in the source is shown as unavailable, not as zero. Categories may not be reported separately in every year.",
            "Change figures compare the first and last selected years only when both endpoint values are available.",
            "The Department notes that 2023/24 pupil-level information was affected by Action Short of Strike; see the source bulletin for details.",
            "Pre-school specialist provision is reported separately from 2024/25; reporting and category definitions vary over time.",
        ],
    }


def main() -> None:
    workbook_url = find_workbook_url()
    workbook_bytes = fetch(workbook_url)
    data = build_data(workbook_bytes, workbook_url)
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    WORKBOOK_PATH.write_bytes(workbook_bytes)
    JSON_PATH.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(f"Downloaded the Department workbook and wrote {len(data['years'])} years to {JSON_PATH}")


if __name__ == "__main__":
    main()
